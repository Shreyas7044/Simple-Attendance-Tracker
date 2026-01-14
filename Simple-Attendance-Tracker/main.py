import openpyxl
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

# Load Excel file
book = openpyxl.load_workbook('attendance.xlsx')
sheet = book.active

r = sheet.max_row
resp = 1

l1 = []
l2 = ""
l3 = []

staff_mails = [
    'staff_ci@gmail.com',
    'staff_python@gmail.com',
    'staff_dm@gmail.com'
]

m1 = "Warning! You can take only one more leave for CI."
m2 = "Warning! You can take only one more leave for Python."
m3 = "Warning! You can take only one more leave for Data Mining."

def savefile():
    book.save('attendance.xlsx')
    print("Attendance saved.")

def mailstu(li, msg):
    from_id = "yourgmail@gmail.com"
    pwd = "your_app_password"

    s = smtplib.SMTP('smtp.gmail.com', 587)
    s.starttls()
    s.login(from_id, pwd)

    for to_id in li:
        message = MIMEMultipart()
        message['Subject'] = 'Attendance Report'
        message.attach(MIMEText(msg, 'plain'))
        s.sendmail(from_id, to_id, message.as_string())

    s.quit()

def mailstaff(mail_id, msg):
    from_id = "yourgmail@gmail.com"
    pwd = "your_app_password"

    s = smtplib.SMTP('smtp.gmail.com', 587)
    s.starttls()
    s.login(from_id, pwd)

    message = MIMEMultipart()
    message['Subject'] = 'Lack of Attendance'
    message.attach(MIMEText(msg, 'plain'))

    s.sendmail(from_id, mail_id, message.as_string())
    s.quit()

def check(no_of_days, row_num, subject):
    global l2, l3

    for i in range(len(row_num)):
        if no_of_days[i] == 2:
            email_id = sheet.cell(row=row_num[i], column=2).value
            l1.append(email_id)
            mailstu(l1, [m1, m2, m3][subject-1])

        elif no_of_days[i] > 2:
            roll = sheet.cell(row=row_num[i], column=1).value
            email_id = sheet.cell(row=row_num[i], column=2).value
            l2 += str(roll) + " "
            l3.append(email_id)

    if l2:
        subject_name = ["CI", "Python", "Data Mining"][subject-1]
        mailstu(l3, f"You have lack of attendance in {subject_name}")
        mailstaff(staff_mails[subject-1], f"Students with shortage: {l2}")

while resp == 1:
    print("1 -> CI\n2 -> Python\n3 -> DM")
    y = int(input("Enter subject: "))
    n = int(input("Number of absentees: "))

    absentees = list(map(int, input("Enter roll numbers: ").split()))
    row_num, no_of_days = [], []

    for roll in absentees:
        for i in range(2, r + 1):
            if sheet.cell(row=i, column=1).value == roll:
                col = y + 2
                count = sheet.cell(row=i, column=col).value + 1
                sheet.cell(row=i, column=col).value = count
                savefile()
                row_num.append(i)
                no_of_days.append(count)

    check(no_of_days, row_num, y)
    resp = int(input("Another subject? (1-Yes, 0-No): "))